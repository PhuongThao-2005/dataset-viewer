#!/usr/bin/env python3
"""
generate_manifest.py
────────────────────
Chạy script này sau khi download dataset từ Kaggle về máy local.
Output: public/manifest.json + ảnh được copy vào public/images/

Usage:
  python generate_manifest.py --dataset /path/to/kaggle/dataset --output ./public

Dataset structure expected (BTRXD / Brain Tumor MRI):
  dataset/
    Training/
      glioma/         ← ảnh gốc
      meningioma/
      pituitary/
      notumor/
    Testing/
      glioma/
      meningioma/
      pituitary/
      notumor/

  (Optional) Nếu có mask riêng:
    dataset_masks/
      Training/glioma/...
      ...

Nếu không có mask sẵn, script sẽ tạo pseudo-mask bằng Otsu threshold.
"""

import os
import json
import shutil
import argparse
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional

# ── Optional: tạo mask bằng OpenCV nếu không có sẵn ──
try:
    import cv2
    import numpy as np
    HAS_CV2 = True
except ImportError:
    HAS_CV2 = False
    print("[WARN] opencv-python không có. Install: pip install opencv-python")
    print("[WARN] Mask sẽ dùng ảnh gốc làm placeholder.")

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

CATEGORY_CONFIG = {
    "glioma":     {"label": "Glioma",      "color": "#ff4444", "description": "Khối u tế bào glial, thường ác tính"},
    "meningioma": {"label": "Meningioma",   "color": "#ffa500", "description": "Khối u màng não, thường lành tính"},
    "pituitary":  {"label": "Pituitary",    "color": "#8b5cf6", "description": "Khối u tuyến yên"},
    "notumor":    {"label": "No Tumor",     "color": "#00ff88", "description": "Ảnh não bình thường"},
    "no_tumor":   {"label": "No Tumor",     "color": "#00ff88", "description": "Ảnh não bình thường"},
    # Thêm categories của dataset bạn vào đây
}

SPLIT_MAP = {
    "Training": "train",
    "training": "train",
    "train": "train",
    "Testing": "test",
    "testing": "test",
    "test": "test",
    "Validation": "val",
    "validation": "val",
    "val": "val",
}

def create_pseudo_mask(img_path: Path, out_path: Path):
    """Tạo binary mask bằng Otsu threshold nếu không có mask sẵn."""
    if not HAS_CV2:
        shutil.copy(img_path, out_path)
        return
    
    img = cv2.imread(str(img_path), cv2.IMREAD_GRAYSCALE)
    if img is None:
        shutil.copy(img_path, out_path)
        return
    
    # Gaussian blur → Otsu threshold
    blurred = cv2.GaussianBlur(img, (5, 5), 0)
    _, mask = cv2.threshold(blurred, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    
    # Morphological operations để làm sạch mask
    kernel = np.ones((3, 3), np.uint8)
    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel, iterations=2)
    mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, kernel, iterations=1)
    
    cv2.imwrite(str(out_path), mask)

def to_web_prefix(prefix: str) -> str:
    cleaned = (prefix or "images").strip().strip("/")
    return cleaned or "images"


def build_web_path(
    category: str,
    image_type: str,
    file_name: str,
    cloud_base_url: str = None,
    cloud_prefix: str = "images",
):
    relative = f"{to_web_prefix(cloud_prefix)}/{category}/{image_type}/{file_name}"
    if cloud_base_url:
        base = cloud_base_url.rstrip("/")
        return f"{base}/{relative}"
    return f"/{relative}"


def safe_name(value: str) -> str:
    return str(value or "").strip()


def normalize_category(value: str) -> str:
    return safe_name(value).lower().replace(" ", "_")


def normalize_split(value: Optional[str], default: str = "train") -> str:
    if value is None:
        return default
    mapped = SPLIT_MAP.get(str(value).strip(), None)
    if mapped is not None:
        return mapped
    lowered = str(value).strip().lower()
    return SPLIT_MAP.get(lowered, default)


def is_positive_flag(value) -> bool:
    text = safe_name(value).lower()
    return text in {"1", "true", "yes", "y"}


def infer_class_from_row(
    row: Dict[str, str],
    class_col: str,
    class_onehot_cols: Optional[List[str]] = None,
) -> str:
    if class_onehot_cols:
        for col in class_onehot_cols:
            if is_positive_flag(row.get(col, "")):
                return col

    fallback = safe_name(row.get(class_col, ""))
    if fallback:
        return fallback
    return "unknown"


def find_mask_file(masks_path: Optional[Path], mask_name: str, stem: str) -> Optional[Path]:
    if masks_path is None:
        return None

    # Priority 1: explicit mask value from xlsx
    if mask_name:
        candidate = masks_path / mask_name
        if candidate.exists():
            return candidate
        fallback = masks_path / Path(mask_name).name
        if fallback.exists():
            return fallback

    # Priority 2: infer by stem (IMG000123 -> IMG000123_mask.* or IMG000123.*)
    patterns = [f"{stem}_mask.*", f"{stem}.*"]
    for pattern in patterns:
        matches = sorted(masks_path.glob(pattern))
        if matches:
            return matches[0]

    return None


def load_rows_from_xlsx(xlsx_path: str) -> List[Dict[str, str]]:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is required for --labels-xlsx. Install: pip install openpyxl")

    wb = load_workbook(filename=xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise RuntimeError("XLSX has no header row")

    headers = [safe_name(h) for h in header_row]
    if not any(headers):
        raise RuntimeError("XLSX header row is empty")

    rows: List[Dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = [safe_name(v) for v in row]
        if not any(values):
            continue
        rows.append({headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))})

    return rows


def generate_manifest_from_labels(
    images_dir: str,
    output_dir: str,
    labels_xlsx: str,
    masks_dir: Optional[str] = None,
    image_col: str = "image_id",
    mask_col: str = "mask",
    class_col: str = "class",
    class_onehot_cols: Optional[List[str]] = None,
    split_col: str = "split",
    default_split: str = "train",
    max_per_category: Optional[int] = None,
    cloud_base_url: Optional[str] = None,
    cloud_prefix: str = "images",
):
    images_path = Path(images_dir)
    output_path = Path(output_dir)
    images_out = output_path / "images"
    masks_path = Path(masks_dir) if masks_dir else None

    rows = load_rows_from_xlsx(labels_xlsx)
    print(f"[INFO] Images: {images_path}")
    print(f"[INFO] Labels: {labels_xlsx}")
    print(f"[INFO] Output:  {output_path}")
    print(f"[INFO] Rows in labels: {len(rows)}")

    category_counter: Dict[str, int] = {}
    categories_found: Dict[str, Dict] = {}
    samples = []
    sample_counter = 0

    for row in rows:
        image_name = safe_name(row.get(image_col, ""))
        if not image_name:
            continue

        class_raw = infer_class_from_row(
            row=row,
            class_col=class_col,
            class_onehot_cols=class_onehot_cols,
        )
        cat_key = normalize_category(class_raw or "unknown")
        split = normalize_split(row.get(split_col), default=default_split)

        if max_per_category and category_counter.get(cat_key, 0) >= max_per_category:
            continue

        img_src = images_path / image_name
        if not img_src.exists():
            img_src = images_path / Path(image_name).name
        if not img_src.exists():
            continue

        stem = img_src.stem
        mask_name = safe_name(row.get(mask_col, ""))
        mask_src = find_mask_file(masks_path=masks_path, mask_name=mask_name, stem=stem)
        orig_name = f"{stem}.jpg"
        mask_file_name = f"{stem}_mask.jpg"

        orig_out_dir = images_out / cat_key / "original"
        mask_out_dir = images_out / cat_key / "mask"
        orig_out_dir.mkdir(parents=True, exist_ok=True)
        mask_out_dir.mkdir(parents=True, exist_ok=True)

        orig_dest = orig_out_dir / orig_name
        mask_dest = mask_out_dir / mask_file_name

        target_size = None
        if HAS_CV2:
            img = cv2.imread(str(img_src), cv2.IMREAD_COLOR)
            if img is not None:
                h, w = img.shape[:2]
                if max(h, w) > 512:
                    scale = 512 / max(h, w)
                    img = cv2.resize(img, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)
                target_size = (img.shape[1], img.shape[0])
                cv2.imwrite(str(orig_dest), img, [cv2.IMWRITE_JPEG_QUALITY, 90])
            else:
                shutil.copy(img_src, orig_dest)
        else:
            shutil.copy(img_src, orig_dest)

        if mask_src is not None:
            if HAS_CV2:
                mimg = cv2.imread(str(mask_src), cv2.IMREAD_COLOR)
                if mimg is not None:
                    if target_size and (mimg.shape[1], mimg.shape[0]) != target_size:
                        mimg = cv2.resize(mimg, target_size, interpolation=cv2.INTER_NEAREST)
                    cv2.imwrite(str(mask_dest), mimg, [cv2.IMWRITE_JPEG_QUALITY, 90])
                else:
                    shutil.copy(mask_src, mask_dest)
            else:
                shutil.copy(mask_src, mask_dest)
        else:
            create_pseudo_mask(img_src, mask_dest)

        cat_config = CATEGORY_CONFIG.get(cat_key, {
            "label": class_raw or cat_key,
            "color": "#64748b",
            "description": f"Category: {class_raw or cat_key}",
        })
        if cat_key not in categories_found:
            categories_found[cat_key] = {"id": cat_key, **cat_config, "count": 0}

        orig_web = build_web_path(cat_key, "original", orig_name, cloud_base_url, cloud_prefix)
        mask_web = build_web_path(cat_key, "mask", mask_file_name, cloud_base_url, cloud_prefix)

        sample_counter += 1
        sample_id = f"image{sample_counter:05d}"
        samples.append({
            "id": sample_id,
            "category": cat_key,
            "split": split,
            "originalPath": orig_web,
            "maskPath": mask_web,
            "metadata": {
                "source": "Kaggle",
                "filename": img_src.name,
            }
        })

        categories_found[cat_key]["count"] += 1
        category_counter[cat_key] = category_counter.get(cat_key, 0) + 1

    manifest = {
        "name": "BTRXD Dataset",
        "description": f"Dataset with {len(samples)} samples, {len(categories_found)} categories",
        "generatedAt": datetime.now().isoformat(),
        "categories": list(categories_found.values()),
        "samples": samples,
    }

    manifest_path = output_path / "manifest.json"
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)

    print(f"\n[OK] Done!")
    print(f"   Samples: {len(samples)}")
    print(f"   Categories: {', '.join(categories_found.keys())}")
    print(f"   Manifest: {manifest_path}")
    print(f"\nNext steps:")
    if cloud_base_url:
        print(f"   1. Upload '{images_out}' to: {cloud_base_url.rstrip('/')}/{to_web_prefix(cloud_prefix)}/")
        print(f"   2. Keep manifest at: {manifest_path}")
        print(f"   3. Run app: npm run dev")
    else:
        print(f"   1. Use generated images in next.js project: public/images/")
        print(f"   2. Keep manifest at: public/manifest.json")
        print(f"   3. Run app: npm run dev")


def generate_manifest(dataset_dir: str, output_dir: str, mask_dir: str = None,
                       max_per_category: int = None, cloud_base_url: str = None,
                       cloud_prefix: str = "images"):
    dataset_path = Path(dataset_dir)
    output_path = Path(output_dir)
    images_out = output_path / "images"
    
    print(f"[INFO] Dataset: {dataset_path}")
    print(f"[INFO] Output:  {output_path}")
    
    samples = []
    categories_found = {}
    
    # Walk the dataset directory
    for split_dir in sorted(dataset_path.iterdir()):
        if not split_dir.is_dir():
            continue
        
        split = SPLIT_MAP.get(split_dir.name)
        if split is None:
            print(f"  [SKIP] {split_dir.name} — unknown split name")
            continue
        
        print(f"\n[INFO] Processing split: {split_dir.name} -> '{split}'")
        
        for cat_dir in sorted(split_dir.iterdir()):
            if not cat_dir.is_dir():
                continue
            
            cat_key = cat_dir.name.lower()
            cat_config = CATEGORY_CONFIG.get(cat_key, {
                "label": cat_dir.name,
                "color": "#64748b",
                "description": f"Category: {cat_dir.name}"
            })
            
            if cat_key not in categories_found:
                categories_found[cat_key] = {
                    "id": cat_key,
                    **cat_config,
                    "count": 0
                }
            
            # Find all images
            img_extensions = {'.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff'}
            images = sorted([f for f in cat_dir.iterdir() 
                            if f.suffix.lower() in img_extensions])
            
            if max_per_category:
                images = images[:max_per_category]
            
            print(f"  [{cat_key}] {len(images)} images")
            
            # Output dirs
            orig_out_dir = images_out / cat_key / "original"
            mask_out_dir = images_out / cat_key / "mask"
            orig_out_dir.mkdir(parents=True, exist_ok=True)
            mask_out_dir.mkdir(parents=True, exist_ok=True)
            
            for img_path in images:
                stem = img_path.stem
                ext = img_path.suffix.lower()
                
                # Output filenames (normalize to .jpg)
                orig_name = f"{stem}.jpg"
                mask_name = f"{stem}_mask.jpg"
                
                orig_dest = orig_out_dir / orig_name
                mask_dest = mask_out_dir / mask_name
                
                # Copy original
                if HAS_CV2:
                    img = cv2.imread(str(img_path))
                    if img is not None:
                        # Resize to max 512px for web performance
                        h, w = img.shape[:2]
                        if max(h, w) > 512:
                            scale = 512 / max(h, w)
                            img = cv2.resize(img, (int(w*scale), int(h*scale)))
                        cv2.imwrite(str(orig_dest), img, [cv2.IMWRITE_JPEG_QUALITY, 90])
                    else:
                        shutil.copy(img_path, orig_dest)
                else:
                    shutil.copy(img_path, orig_dest)
                
                # Handle mask
                if mask_dir:
                    # Look for corresponding mask
                    mask_src = Path(mask_dir) / split_dir.name / cat_dir.name / img_path.name
                    if not mask_src.exists():
                        # Try _mask suffix
                        mask_src = Path(mask_dir) / split_dir.name / cat_dir.name / f"{stem}_mask{ext}"
                    
                    if mask_src.exists():
                        shutil.copy(mask_src, mask_dest)
                    else:
                        create_pseudo_mask(img_path, mask_dest)
                else:
                    create_pseudo_mask(img_path, mask_dest)
                
                # Build web paths (local /images/... or absolute cloud URLs)
                orig_web = build_web_path(
                    category=cat_key,
                    image_type="original",
                    file_name=orig_name,
                    cloud_base_url=cloud_base_url,
                    cloud_prefix=cloud_prefix,
                )
                mask_web = build_web_path(
                    category=cat_key,
                    image_type="mask",
                    file_name=mask_name,
                    cloud_base_url=cloud_base_url,
                    cloud_prefix=cloud_prefix,
                )
                
                samples.append({
                    "id": f"{cat_key}_{split}_{stem}",
                    "category": cat_key,
                    "split": split,
                    "originalPath": orig_web,
                    "maskPath": mask_web,
                    "metadata": {
                        "source": "Kaggle",
                        "filename": img_path.name,
                    }
                })
                
                categories_found[cat_key]["count"] = \
                    categories_found[cat_key].get("count", 0) + 1
    
    # Build manifest
    manifest = {
        "name": "BTRXD Dataset",
        "description": f"Dataset với {len(samples)} samples, {len(categories_found)} categories",
        "generatedAt": datetime.now().isoformat(),
        "categories": list(categories_found.values()),
        "samples": samples,
    }
    
    manifest_path = output_path / "manifest.json"
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    
    print(f"\n Done!")
    print(f"   Samples: {len(samples)}")
    print(f"   Categories: {', '.join(categories_found.keys())}")
    print(f"   Manifest: {manifest_path}")
    print(f"\nNext steps:")
    if cloud_base_url:
        print(f"   1. Upload '{images_out}' to: {cloud_base_url.rstrip('/')}/{to_web_prefix(cloud_prefix)}/")
        print(f"   2. Keep manifest at: {manifest_path}")
        print(f"   3. Run app: npm run dev")
    else:
        print(f"   1. Use generated images in next.js project: public/images/")
        print(f"   2. Keep manifest at: public/manifest.json")
        print(f"   3. Run app: npm run dev")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate manifest for BTRXD viewer")
    parser.add_argument("--dataset", required=True, help="Path to Kaggle dataset folder")
    parser.add_argument("--images-dir", default=None, help="Flat images folder path (used with --labels-xlsx)")
    parser.add_argument("--output", default="./public", help="Output directory (default: ./public)")
    parser.add_argument("--masks", default=None, help="Optional separate mask folder")
    parser.add_argument("--labels-xlsx", default=None, help="XLSX labels file for flat image/mask folders")
    parser.add_argument("--image-col", default="image_id", help="Column name for image file in labels xlsx")
    parser.add_argument("--mask-col", default="mask", help="Column name for mask file in labels xlsx")
    parser.add_argument("--class-col", default="class", help="Column name for class label in labels xlsx")
    parser.add_argument("--class-onehot-cols", default=None,
                        help="Comma-separated one-hot class columns, first positive will be selected")
    parser.add_argument("--split-col", default="split", help="Column name for split label in labels xlsx")
    parser.add_argument("--default-split", default="train", help="Fallback split if xlsx has no split value")
    parser.add_argument("--max", type=int, default=None,
                        help="Max images per category (useful for testing, e.g. --max 20)")
    parser.add_argument("--cloud-base-url", default=None,
                        help="Optional cloud storage public base URL. Example: https://cdn.example.com")
    parser.add_argument("--cloud-prefix", default="images",
                        help="URL path prefix under cloud base URL (default: images)")
    
    args = parser.parse_args()
    if args.labels_xlsx:
        input_images = args.images_dir or args.dataset
        generate_manifest_from_labels(
            images_dir=input_images,
            output_dir=args.output,
            labels_xlsx=args.labels_xlsx,
            masks_dir=args.masks,
            image_col=args.image_col,
            mask_col=args.mask_col,
            class_col=args.class_col,
            class_onehot_cols=[c.strip() for c in args.class_onehot_cols.split(",")] if args.class_onehot_cols else None,
            split_col=args.split_col,
            default_split=args.default_split,
            max_per_category=args.max,
            cloud_base_url=args.cloud_base_url,
            cloud_prefix=args.cloud_prefix,
        )
    else:
        generate_manifest(
            dataset_dir=args.dataset,
            output_dir=args.output,
            mask_dir=args.masks,
            max_per_category=args.max,
            cloud_base_url=args.cloud_base_url,
            cloud_prefix=args.cloud_prefix,
        )
